"""Z023 consolidado: lectura tipada, aportes por meta (metas compartidas) y validaciones. Usa el ejemplo ficticio."""
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook

from src.evaplan import aportes as A
from src.evaplan import lectura as L
from src.evaplan import limpieza as lz
from src.evaplan import pipeline
from src.evaplan import seguimiento as S
from src.evaplan import validaciones as V

EJ = Path(__file__).resolve().parent.parent / "data" / "ejemplos"
PI, CE, DR, Z = (EJ / "ejemplo_PI_MP_evaplan.xlsx", EJ / "ejemplo_Centralizadas.xlsx",
                 EJ / "ejemplo_PI_Drive.xlsx", EJ / "ejemplo_Z023.xlsx")
MP_COMPARTIDA = "MP9900101019901001"


@pytest.fixture(scope="module")
def z():
    return L.leer_z023(Z)


@pytest.fixture(scope="module")
def res():
    return pipeline.ejecutar(PI, CE, DR, z023=Z)


def test_lectura_tipos_y_llave(z):
    assert len(z) == 9 and z["ppm_actividad"].is_unique
    assert str(z["vigencia"].dtype) == "Int64" and set(z["vigencia"]) == {2025, 2026}
    assert z["valor_actividad"].sum() == 6_200_000_000                       # suma de la columna numérica
    assert z["codigo_mr"].iloc[0] == "99001"                                 # 'MR99001' -> '99001'
    assert z["centro_gestor_mp"].isna().all()                                 # '#ERROR!' de fórmula = vacío


def test_codigo_proyecto_ps_sale_del_codigo_de_actividad(z):
    fila = z[z["ppm_actividad"] == "0000000001"].iloc[0]
    assert fila["codigo_proyecto_ps"] == "PI99-000001" and fila["proyecto_ppm"] == "PI-900001"
    descentralizada = z[z["dependencia"] == "0099"].iloc[0]
    assert pd.isna(descentralizada["ps_actividad"]) and pd.isna(descentralizada["codigo_proyecto_ps"])


def test_error_claro_si_no_es_un_z023(tmp_path):
    wb = Workbook()
    wb.active.title = "Hoja1"
    wb.active.append(["Otra", "cosa"])
    ruta = tmp_path / "otro.xlsx"
    wb.save(ruta)
    with pytest.raises(L.EsquemaError, match="Z023"):
        L.leer_z023(ruta)


def test_valor_acepta_texto_con_comas_de_miles_y_errores_de_formula():
    assert lz.moneda(" $  3,200,000 ") == 3_200_000
    assert lz.moneda("2.339.400.000") == 2_339_400_000          # el formato es-CO sigue igual
    assert lz.moneda("1,5") == 1.5
    assert lz.texto("#ERROR!") is pd.NA and lz.codigo("#N/A") is pd.NA


def test_aportes_separan_lo_propio_de_lo_ajeno(res):
    a = res.aportes
    assert set(a["vigencia"]) == {2026}                                      # ignora la fila de 2025
    meta = a[a["codigo_mp"] == MP_COMPARTIDA].set_index("codigo_dependencia")
    assert list(meta.index) == ["9999", "0099", "9998"]                      # propia primero
    assert meta.loc["9999", "es_propia"] and not meta.loc["9998", "es_propia"]
    assert meta.loc["0099", "tipo_entidad"] == "Descentralizada" and meta.loc["0099", "n_sin_codigo_ps"] == 1
    assert meta.loc["9999", "n_actividades"] == 2 and meta.loc["9999", "valor_actividades"] == 2_500_000_000


def test_matriz_incluye_columnas_de_z023_solo_si_se_carga(res):
    assert set(S.COLUMNAS_Z023) <= set(res.matriz.columns)
    sin = pipeline.ejecutar(PI, CE, DR)
    assert not set(S.COLUMNAS_Z023) & set(sin.matriz.columns) and not sin.usa_z023 and sin.aportes.empty
    f = res.matriz.set_index("codigo_mp").loc[MP_COMPARTIDA]
    assert (f["n_proyectos_z023"], f["n_entidades_aportantes"], f["n_proyectos_ajenos"]) == (3, 3, 2)
    assert "9998" in f["aportes_otras_entidades"] and "0099" in f["aportes_otras_entidades"]


def test_hallazgos_de_z023_son_informativos(res):
    h = res.hallazgos[res.hallazgos["regla"].isin(["meta_con_aportes_de_otras_entidades", "meta_sin_proyectos_en_z023"])]
    assert set(h["severidad"]) == {"info"}
    assert (h["regla"] == "meta_con_aportes_de_otras_entidades").sum() == 1
    assert set(h.loc[h["regla"] == "meta_sin_proyectos_en_z023", "llave"]) == {"MP9900202019902001", "MP9900202039902003"}
    # lo informativo no suma a "con alertas": la meta compartida no tiene alertas por esto
    assert res.matriz.set_index("codigo_mp").loc[MP_COMPARTIDA, "alertas"].count("[Informativa]") >= 1


def test_actividad_de_centralizadas_ausente_del_z023(res):
    q = res.calidad
    assert list(q.loc[q["regla"] == "actividad_no_esta_en_z023", "llave"]) == ["PI99-000002/1/1/01/02"]
    assert "mp_distinta_en_z023" not in set(q["regla"])


def test_calidad_propia_del_z023(z):
    q = V.z023_calidad(z)
    assert set(q["regla"]) == {"z023_sin_meta_producto", "z023_bpin_no_valido"}
    assert list(q.loc[q["regla"] == "z023_bpin_no_valido", "llave"]) == ["PI-900005"]
    assert (q["severidad"] == "advertencia").all()


def test_calidad_del_z023_solo_revisa_las_entidades_cargadas(res):
    # Las anomalías sembradas son de la dependencia 9998, que no está en los archivos cargados: no se reportan.
    assert not {"z023_sin_meta_producto", "z023_bpin_no_valido"} & set(res.calidad["regla"])


def test_excel_incluye_hoja_de_aportes_solo_con_z023(res):
    import io

    from openpyxl import load_workbook

    from src.evaplan import reportes
    con = load_workbook(io.BytesIO(reportes.a_excel(res.matriz, res.hallazgos, res.calidad, res.aportes)))
    assert "Aportes_Z023" in con.sheetnames
    sin = load_workbook(io.BytesIO(reportes.a_excel(res.matriz, res.hallazgos, res.calidad)))
    assert "Aportes_Z023" not in sin.sheetnames


def test_resultado_informa_las_filas_del_z023_por_vigencia(res):
    assert res.z023_filas_por_vigencia == {2025: 1, 2026: 8}
    assert pipeline.ejecutar(PI, CE, DR).z023_filas_por_vigencia == {}
