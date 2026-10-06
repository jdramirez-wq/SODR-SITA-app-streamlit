import io
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.evaplan import lectura as L
from src.evaplan import pipeline, reportes
from src.evaplan import seguimiento as S
from src.evaplan import validaciones as V

EJ = Path(__file__).resolve().parent.parent / "data" / "ejemplos"
PI, CE, DR = (EJ / "ejemplo_PI_MP_evaplan.xlsx", EJ / "ejemplo_Centralizadas.xlsx", EJ / "ejemplo_PI_Drive.xlsx")


@pytest.fixture(scope="module")
def datos():
    return L.leer_pi_mp_evaplan(PI), L.leer_centralizadas(CE), L.leer_pi_drive_mp(DR)


@pytest.fixture(scope="module")
def matriz(datos):
    return S.construir_matriz(*datos).set_index("codigo_mp")


def test_universo_incluye_metas_sin_reporte(matriz):
    assert len(matriz) == 5 and (matriz["estado_reporte"] == "Reportada").sum() == 4
    assert matriz.loc["MP9900202039902003", "estado_reporte"] == "Sin reporte en EVAPLAN"
    assert matriz["vigencia"].eq(2026).all()                       # inferida de los encabezados 'VAL ALC'


def test_sin_drive_solo_hay_metas_reportadas(datos):
    pi, ce, _ = datos
    m = S.construir_matriz(pi, ce, None, vigencia=2026)
    assert len(m) == 4 and (m["estado_reporte"] == "Reportada").all()


def test_avance_acumulado_usa_logros_previos_y_resultado(matriz):
    f = matriz.loc["MP9900101019901001"]                           # Acumulado: PG 100; logros 20+30; meta 2026 = 30
    assert f["logro_previo"] == 50 and f["resultado"] == 12
    assert f["pct_avance_vigencia"] == pytest.approx(12 / 30)
    assert f["avance_cuatrienio"] == 62 and f["pct_avance_pg"] == pytest.approx(0.62)


def test_flujo_no_acumula_logros_previos(matriz):
    f = matriz.loc["MP9900101029901002"]
    assert pd.isna(f["logro_previo"]) and pd.isna(f["pct_avance_pg"])
    assert f["pct_avance_vigencia"] == pytest.approx(41.667 / 85)


def test_presupuesto_y_avance_de_actividades_por_proyecto(matriz):
    f = matriz.loc["MP9900101019901001"]
    assert f["ppto_definitivo"] == 2_500_000_000 and f["ppto_obligaciones"] == 1_500_000_000
    assert f["pct_ejecucion_financiera"] == pytest.approx(0.6)
    assert f["avance_actividades"] == pytest.approx(0.125)         # (25 % + 0 %) / 2
    assert f["avance_por_proyecto"] == "PI99-000001: 12.5 %"
    assert f["n_proyectos"] == 1 and f["n_registros"] == 2
    assert not matriz.loc["MP9900202019902001", "tiene_plan_de_accion"]


def test_el_promedio_se_calcula_por_proyecto_cuando_hay_varios(datos):
    pi, ce, dr = datos
    ce = ce.copy()
    ce.loc[ce["codigo_actividad"] == "PI99-000001/1/2/01/03", ["codigo_mp", "codigo_proyecto", "avance_actividad_pct"]] = \
        ["MP9900101019901001", "PI99-000009", 80.0]
    f = S.construir_matriz(pi, ce, dr).set_index("codigo_mp").loc["MP9900101019901001"]
    assert f["n_proyectos"] == 2 and "PI99-000001: 12.5 %" in f["avance_por_proyecto"] \
        and "PI99-000009: 80.0 %" in f["avance_por_proyecto"]


def test_focalizacion_replica_la_logica_original(datos):
    pi, ce, dr = datos
    f = S.construir_matriz(pi, ce, dr).set_index("codigo_mp").loc["MP9900101019901001"]
    assert f["focalizacion"] == "¿Cuál Otro?: Servidores públicos"       # Otros=0 pero hay texto, como en la página original


def test_hallazgos_sembrados(matriz, datos):
    h = S.detectar_hallazgos(matriz.reset_index())
    pares = set(zip(h["regla"], h["llave"]))
    assert ("sin_reporte", "MP9900202039902003") in pares
    assert ("sin_plan_de_accion", "MP9900202019902001") in pares
    assert ("registros_con_obligaciones_sin_avance", "MP9900202029902002") in pares
    # Criterio por defecto: el avance 0 se justifica en Dificultades (esta meta la tiene) y se avisa, como
    # informativo, que además trae logro/análisis.
    assert ("narrativa_con_resultado_cero", "MP9900202019902001") in pares
    assert not {r for r, _ in pares} & {"avance_cero_sin_justificacion", "ejecucion_sin_avance_sin_explicacion"}


def test_avance_sin_obligaciones_exige_justificar_gestion(datos):
    pi, ce, dr = datos
    ce0 = ce.copy()
    ce0.loc[ce0["codigo_mp"] == "MP9900101029901002", "ppto_obligaciones"] = 0.0
    h = S.detectar_hallazgos(S.construir_matriz(pi, ce0, dr))
    assert ("avance_sin_ejecucion_financiera", "advertencia") in set(zip(h["regla"], h["severidad"]))
    pi2 = pi.copy()
    pi2.loc[pi2["codigo_mp"] == "MP9900101029901002", "principal_logro"] = "Se logró por gestión con otra entidad"
    h2 = S.detectar_hallazgos(S.construir_matriz(pi2, ce0, dr))
    assert ("avance_sin_ejecucion_con_gestion", "info") in set(zip(h2["regla"], h2["severidad"]))
    assert "avance_sin_ejecucion_financiera" not in set(h2["regla"])


def test_no_se_inventan_umbrales(matriz):
    """Una meta con 100 % de ejecución financiera y 0 % de avance físico NO genera juicio de 'crítico' por umbral:
    solo hechos. La única alerta de ese estilo es la objetiva (obligaciones > 0 sin avance de actividades)."""
    reglas = set(S.detectar_hallazgos(matriz.reset_index())["regla"])
    assert not any("umbral" in r or "critic" in r or "semaforo" in r for r in reglas)


def test_pipeline_pasa_el_criterio_flexible():
    est, fle = pipeline.ejecutar(PI, CE, DR), pipeline.ejecutar(PI, CE, DR, criterio_flexible=True)
    assert "narrativa_con_resultado_cero" in set(est.hallazgos["regla"])
    assert "narrativa_con_resultado_cero" not in set(fle.hallazgos["regla"])


def test_pipeline_completo_y_reportes():
    r = pipeline.ejecutar(PI, CE, DR)
    assert r.vigencia == 2026 and r.usa_drive and r.resumen["metas_sin_reporte"] == 1
    assert not r.calidad.empty and (r.calidad["severidad"] != "error").all()
    wb = load_workbook(io.BytesIO(reportes.a_excel(r.matriz, r.hallazgos, r.calidad)))
    assert wb.sheetnames == ["MP_PI_PA", "Hallazgos", "Sin_reporte", "Calidad_de_datos"]
    assert wb["MP_PI_PA"].max_row == 6
    pdf = reportes.a_pdf(r.matriz)
    assert pdf.startswith(b"%PDF") and len(pdf) > 3000


def test_pipeline_sin_drive_y_con_drive_caido():
    r = pipeline.ejecutar(PI, CE, None, vigencia=2026)
    assert not r.usa_drive and r.resumen["metas_sin_reporte"] == 0
    r2 = pipeline.ejecutar(PI, CE, "https://no-existe.invalid/x.xlsx", vigencia=2026)
    assert r2.avisos and "Drive" in r2.avisos[0] and len(r2.matriz) == 4


def test_pipeline_con_archivo_subido_se_puede_leer_dos_veces():
    pi, ce = io.BytesIO(PI.read_bytes()), io.BytesIO(CE.read_bytes())
    pipeline.ejecutar(pi, ce, None, vigencia=2026)
    r = pipeline.ejecutar(pi, ce, None, vigencia=2026)            # segundo clic en "Procesar"
    assert len(r.matriz) == 4


def test_estructura_incorrecta_da_error_claro():
    with pytest.raises(L.EsquemaError, match="Código de Meta"):
        pipeline.ejecutar(CE, CE, None)                           # el usuario subió Centralizadas en el lugar del PI


def test_valor_proyectado_es_proyeccion_de_cierre(datos):
    pi, ce, dr = datos
    m = S.construir_matriz(pi, ce, dr).set_index("codigo_mp")
    f = m.loc["MP9900101019901001"]                                # proyectado 25 vs meta 30
    assert f["valor_proyectado"] == 25 and f["pct_proyectado_vs_meta"] == pytest.approx(25 / 30)
    assert pd.isna(m.loc["MP9900101029901002", "pct_proyectado_vs_meta"])    # sin proyección: no se inventa
    h = S.detectar_hallazgos(m.reset_index())
    assert ("proyeccion_bajo_meta", "MP9900101019901001") in set(zip(h["regla"], h["llave"]))


def test_proyeccion_menor_que_el_resultado_es_incoherente_en_acumulados(datos):
    pi, ce, dr = datos
    pi = pi.copy()
    pi.loc[pi["codigo_mp"] == "MP9900101019901001", "valor_proyectado"] = 5.0     # resultado ya es 12
    h = S.detectar_hallazgos(S.construir_matriz(pi, ce, dr))
    assert ("proyeccion_menor_que_resultado", "MP9900101019901001") in set(zip(h["regla"], h["llave"]))


def test_no_se_agrupa_por_actividad_se_cuentan_registros(matriz):
    """El código de actividad no es confiable para agrupar: la unidad es el registro presupuestal."""
    f = matriz.loc["MP9900202029902002"]
    assert f["n_registros"] == 3 and "n_actividades" not in matriz.columns
    assert f["ppto_definitivo"] == 1_100_000_000                    # los presupuestos de los registros se suman


def _hallazgos(datos, ce=None, pi=None):
    p, c, d = datos
    return S.detectar_hallazgos(S.construir_matriz(p if pi is None else pi, c if ce is None else ce, d))


def test_obligaciones_sin_avance_bajan_a_info_si_la_observacion_lo_explica(datos):
    _, ce, _ = datos
    ce = ce.copy()
    ce.loc[ce["codigo_mp"] == "MP9900202029902002", "observacion"] = "El entregable está para el mes de noviembre."
    h = _hallazgos(datos, ce=ce)
    reglas = set(zip(h["regla"], h["severidad"], h["llave"]))
    assert ("registros_sin_avance_con_observacion", "info", "MP9900202029902002") in reglas
    assert "registros_con_obligaciones_sin_avance" not in set(h["regla"])


def test_registro_sin_cantidad_programada_no_genera_alerta_de_avance(datos):
    from src.evaplan import validaciones as V
    _, ce, _ = datos
    ce = ce.copy()
    ce["cant_programada_vigencia"] = pd.NA
    assert "financiero_sin_fisico" not in set(V.ejecucion_financiera_vs_fisica(ce)["regla"])


@pytest.mark.parametrize("analisis, dificultades, flexible, alerta", [
    (None, None, False, True),                        # sin ninguna justificación
    ("Se explica en el análisis", None, False, True),     # por defecto debe estar en Dificultades
    ("Se explica en el análisis", None, True, False),     # criterio flexible: el análisis también sirve
    (None, "Etapa precontractual", False, False),         # Dificultades siempre sirve
    (None, "Etapa precontractual", True, False),
    (None, None, True, True),
])
def test_resultado_cero_se_justifica_en_dificultades_y_opcionalmente_en_analisis(datos, analisis, dificultades, flexible, alerta):
    pi, ce, dr = datos
    pi = pi.copy()
    i = pi.index[pi["codigo_mp"] == "MP9900202019902001"][0]          # resultado 0, sin plan de acción
    pi.loc[i, ["analisis_logro", "dificultades_gestiones", "principal_logro"]] = [analisis, dificultades, pd.NA]
    h = S.detectar_hallazgos(S.construir_matriz(pi, ce, dr, criterio_flexible=flexible), flexible)
    assert ("avance_cero_sin_justificacion" in set(h["regla"])) is alerta


def test_la_nota_de_logro_con_resultado_cero_solo_aplica_en_el_criterio_estricto(datos):
    pi, ce, dr = datos
    estricto = S.detectar_hallazgos(S.construir_matriz(pi, ce, dr))
    flexible = S.detectar_hallazgos(S.construir_matriz(pi, ce, dr, criterio_flexible=True), True)
    assert "narrativa_con_resultado_cero" in set(estricto["regla"])
    assert "narrativa_con_resultado_cero" not in set(flexible["regla"])


def test_resultado_cero_con_obligaciones_y_sin_justificacion_usa_la_regla_especifica(datos):
    pi, ce, dr = datos
    pi = pi.copy()
    i = pi.index[pi["codigo_mp"] == "MP9900101019901001"][0]
    pi.loc[i, ["resultado", "analisis_logro", "dificultades_gestiones", "principal_logro"]] = [0.0, "Análisis", pd.NA, pd.NA]
    h = S.detectar_hallazgos(S.construir_matriz(pi, ce, dr))
    reglas = set(h[h["llave"] == "MP9900101019901001"]["regla"])
    assert "ejecucion_sin_avance_sin_explicacion" in reglas and "avance_cero_sin_justificacion" not in reglas


def test_si_la_meta_del_export_difiere_de_drive_se_usa_drive_y_se_avisa(datos):
    pi, ce, dr = datos
    pi = pi.copy()
    pi.loc[pi["codigo_mp"] == "MP9900101019901001", "valor_2026"] = 20.0          # Drive dice 30
    m = S.construir_matriz(pi, ce, dr).set_index("codigo_mp")
    f = m.loc["MP9900101019901001"]
    assert f["meta_vigencia"] == 30 and f["meta_vigencia_export"] == 20
    assert f["pct_avance_vigencia"] == pytest.approx(12 / 30)                     # calcula con la de Drive
    h = S.detectar_hallazgos(m.reset_index())
    d = h[h["regla"] == "meta_vigencia_difiere_del_export"]
    assert d["llave"].tolist() == ["MP9900101019901001"] and "20" in d.iloc[0]["detalle"] and "30" in d.iloc[0]["detalle"]


def test_toda_regla_emitida_tiene_nombre_legible(datos):
    from src.evaplan import validaciones as V
    from src.evaplan.reglas import ETIQUETAS_REGLAS, etiquetar
    pi, ce, dr = datos
    r = pipeline.ejecutar(PI, CE, DR)
    emitidas = set(r.hallazgos["regla"]) | set(r.calidad["regla"])
    assert emitidas <= set(ETIQUETAS_REGLAS), emitidas - set(ETIQUETAS_REGLAS)
    assert "hallazgo" in etiquetar(r.hallazgos).columns and V is not None


def test_avance_vacio_con_cantidad_programada_cuenta_como_cero_en_el_promedio():
    ce = pd.DataFrame({
        "codigo_mp": ["MP1"] * 4, "codigo_proyecto": ["P1", "P1", "P1", "P2"], "nombre_proyecto": ["a"] * 4,
        "avance_actividad_pct": [100.0, pd.NA, pd.NA, pd.NA], "cant_programada_vigencia": [1.0, 5.0, 0.0, pd.NA],
        "ppto_definitivo": [10.0] * 4, "ppto_obligaciones": [0.0] * 4, "ppto_disponible": [10.0] * 4,
        "observacion": [pd.NA] * 4, "fila_excel": [3, 4, 5, 6]})
    p = S._consolidar_plan_de_accion(ce).iloc[0]
    # P1: 100 y (vacío con programada 5 -> 0); el vacío con programada 0 no cuenta -> (100+0)/2 = 50 %; P2 sin dato
    assert p["avance_actividades"] == 0.5
    assert p["avance_por_proyecto"] == "P1: 50.0 %"
    q = V.avance_vacio_con_programacion(ce)
    assert list(q["llave"]) == ["MP1"] and "1 de 4" in q.loc[0, "detalle"] and q.loc[0, "severidad"] == "info"


def test_dos_promedios_de_avance_prima_el_de_todas_las_actividades():
    ce = pd.DataFrame({
        "codigo_mp": ["MP1"] * 4, "codigo_proyecto": ["P1"] * 4, "nombre_proyecto": ["a"] * 4,
        "avance_actividad_pct": [80.0, 40.0, 0.0, 0.0], "cant_programada_vigencia": [1.0] * 4,
        "ppto_definitivo": [10.0] * 4, "ppto_obligaciones": [5.0, 1.0, 0.0, pd.NA], "ppto_disponible": [5.0] * 4,
        "observacion": [pd.NA] * 4, "fila_excel": [3, 4, 5, 6]})
    p = S._consolidar_plan_de_accion(ce).iloc[0]
    assert p["avance_actividades"] == pytest.approx(0.30)                   # (80+40+0+0)/4: TODAS
    assert p["avance_actividades_con_obligaciones"] == pytest.approx(0.60)  # (80+40)/2: solo con obligaciones > 0
    assert p["n_registros_con_obligaciones"] == 2 and S.PROMEDIO_QUE_PRIMA == "todas"


def test_sin_obligaciones_el_promedio_complementario_es_sin_dato_no_cero():
    ce = pd.DataFrame({
        "codigo_mp": ["MP1"] * 2, "codigo_proyecto": ["P1"] * 2, "nombre_proyecto": ["a"] * 2,
        "avance_actividad_pct": [10.0, 20.0], "cant_programada_vigencia": [1.0] * 2,
        "ppto_definitivo": [10.0] * 2, "ppto_obligaciones": [0.0, pd.NA], "ppto_disponible": [10.0] * 2,
        "observacion": [pd.NA] * 2, "fila_excel": [3, 4]})
    p = S._consolidar_plan_de_accion(ce).iloc[0]
    assert pd.isna(p["avance_actividades_con_obligaciones"]) and p["avance_actividades"] == pytest.approx(0.15)
    vacio = S._consolidar_plan_de_accion(ce.assign(ppto_obligaciones=pd.NA)).iloc[0]
    assert pd.isna(vacio["ppto_obligaciones"])                              # todo vacío: "sin dato", no 0


def test_el_prompt_explica_cual_promedio_prima():
    from src.evaplan.prompts import generar_prompt_sistema
    t = generar_prompt_sistema("Revisión a Cierre de Vigencia", 2026, True)
    assert "El que PRIMA para tu análisis" in t and "TODAS las actividades" in t
    assert "PRIMA" not in generar_prompt_sistema("Revisión a Cierre de Vigencia", 2026, False)
