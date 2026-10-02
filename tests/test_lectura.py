from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.evaplan import lectura as L

EJ = Path(__file__).resolve().parent.parent / "data" / "ejemplos"


@pytest.fixture(scope="module")
def pi():
    return L.leer_pi_mp_evaplan(EJ / "ejemplo_PI_MP_evaplan.xlsx")


@pytest.fixture(scope="module")
def ce():
    return L.leer_centralizadas(EJ / "ejemplo_Centralizadas.xlsx")


@pytest.fixture(scope="module")
def drive_mp():
    return L.leer_pi_drive_mp(EJ / "ejemplo_PI_Drive.xlsx")


@pytest.fixture(scope="module")
def drive_mr():
    return L.leer_pi_drive_mr(EJ / "ejemplo_PI_Drive.xlsx")


def test_pi_mp_evaplan_estructura_y_tipos(pi):
    assert len(pi) == 4 and pi["codigo_mp"].is_unique
    assert str(pi["codigo_mp"].dtype) == "string" and str(pi["valor_2025"].dtype) == "Float64"
    assert pi["codigo_mp_valido"].all()
    assert set(pi["programa_codigo"]) == {"99"}
    assert pi["fila_excel"].tolist() == [3, 4, 5, 6]          # fila real en Excel (título + encabezado + datos)


def test_np_se_separa_del_cero(pi):
    fila = pi[pi["codigo_mp"] == "MP9900202029902002"].iloc[0]
    assert fila["valor_2024_np"] and pd.isna(fila["valor_2024"])
    assert fila["valor_2025"] == 4 and not fila["valor_2025_np"]


def test_encabezados_numericos_de_anio_se_normalizan(pi):
    assert {"valor_2024", "valor_2025", "valor_2026", "valor_2027"} <= set(pi.columns)


def test_marcadores_vacio_y_columnas_ignoradas(pi):
    assert pi["principal_logro"].isna().sum() == 1             # la meta con resultado 0 no trae logro
    assert "Indígena" not in pi.columns and "Cali" not in pi.columns


def test_centralizadas_dinero_y_codigos(ce):
    assert len(ce) == 5
    fila = ce.iloc[0]
    assert fila["ppto_inicial"] == 2_000_000_000 and fila["ppto_obligaciones"] == 1_500_000_000
    assert ce.loc[1, "ppto_obligaciones"] == 0                  # entero 0 mezclado con texto
    assert fila["bpin"] == "2024009990001" and fila["codigo_producto_mga"] == "9901001"
    assert ce["observacion"].isna().sum() == 4                  # '.' = vacío
    assert ce["ppto_gestion"].isna().all()
    assert (ce["producto_mga_mp"] == ce["codigo_producto_mga"]).all()


def test_drive_mp_lee_por_posicion_y_normaliza(drive_mp):
    assert len(drive_mp) == 6                                   # filas vacías finales descartadas
    fila = drive_mp[drive_mp["codigo_mp"] == "MP9900202039902003"].iloc[0]
    assert fila["entidad_codigo"] == "9999" and fila["entidad_nombre"] == "SECRETARÍA DE EJEMPLO"
    assert fila["comportamiento"] == "Incremento Flujo"        # venía como 'Incremento flujo'
    assert pd.isna(fila["linea_base"]) and pd.isna(fila["anio_linea_base"])   # 'NO DISPONIBLE', '2020-2023'
    assert fila["fecha_actualizacion_indicador"] == pd.Timestamp("2024-12-12")
    assert drive_mp["edt"].dtype == "boolean"
    assert drive_mp["recurso_total_2026"].notna().all()
    assert drive_mp["valor_2024_np"].sum() == 1                 # una 'NP' en la vigente 2024


def test_drive_mr_corrige_codigo_duplicado(drive_mr):
    assert drive_mr["codigo_mr"].tolist() == ["99001", "99002", "98001"]
    assert drive_mr.loc[0, "meta_resultado_nombre"].startswith("ALCANZAR 40 PUNTOS")
    assert drive_mr.loc[2, "comportamiento"] == "Reducción Anual"


def test_error_claro_si_cambia_la_estructura(tmp_path):
    wb = load_workbook(EJ / "ejemplo_Centralizadas.xlsx")
    ws = wb.active
    ws["G2"] = "Código MP"                                      # renombrado: antes 'Cód. MP'
    ruta = tmp_path / "cambiado.xlsx"
    wb.save(ruta)
    with pytest.raises(L.EsquemaError, match="falta la columna 'Cód. MP'"):
        L.leer_centralizadas(ruta)


def test_error_si_se_desplaza_una_columna_de_drive(tmp_path):
    wb = load_workbook(EJ / "ejemplo_PI_Drive.xlsx")
    wb["MP"].insert_cols(3)                                     # corre todo una posición
    ruta = tmp_path / "desplazado.xlsx"
    wb.save(ruta)
    with pytest.raises(L.EsquemaError, match="se esperaba 'Entidad'"):
        L.leer_pi_drive_mp(ruta)


def test_error_si_falta_la_hoja(tmp_path):
    wb = load_workbook(EJ / "ejemplo_PI_Drive.xlsx")
    del wb["MR"]
    ruta = tmp_path / "sin_mr.xlsx"
    wb.save(ruta)
    with pytest.raises(L.EsquemaError, match="hoja 'MR'"):
        L.leer_pi_drive_mr(ruta)


def test_drive_acepta_que_el_tecnico_renombre_el_ano_cerrado(tmp_path):
    """Al cerrar 2026 el técnico cambia el encabezado '2026' por 'VAL ALC 2026': el lector no debe romperse."""
    wb = load_workbook(EJ / "ejemplo_PI_Drive.xlsx")
    wb["MP"]["AM2"] = "VAL ALC 2026"                              # columna AM = posición 39 (valor_2026)
    ruta = tmp_path / "cerrado_2026.xlsx"
    wb.save(ruta)
    df = L.leer_pi_drive_mp(ruta)
    assert df["valor_2026_logro"].all() and not df["valor_2027_logro"].any()
    from src.evaplan import validaciones as V
    assert V.inferir_vigencia(df) == 2027


def test_drive_marca_de_logro_por_ano(drive_mp, drive_mr):
    assert drive_mp[["valor_2024_logro", "valor_2025_logro", "valor_2026_logro", "valor_2027_logro"]].iloc[0].tolist() \
        == [True, True, False, False]
    assert drive_mr["valor_2025_logro"].iloc[0] is not None and not drive_mr["valor_2025_logro"].any()
