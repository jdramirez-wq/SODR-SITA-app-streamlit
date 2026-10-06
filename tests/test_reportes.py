"""Los reportes deben poder leerse sin adivinar: un solo formato de cifras y nada de fracciones ni emojis."""
import io
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from src.evaplan import formato as F
from src.evaplan import pipeline, reportes
from src.evaplan.prompts import generar_prompt_sistema

EJ = Path(__file__).resolve().parent.parent / "data" / "ejemplos"


@pytest.fixture(scope="module")
def res():
    return pipeline.ejecutar(EJ / "ejemplo_PI_MP_evaplan.xlsx", EJ / "ejemplo_Centralizadas.xlsx",
                             EJ / "ejemplo_PI_Drive.xlsx", z023=EJ / "ejemplo_Z023.xlsx")


@pytest.fixture(scope="module")
def texto_pdf(res):
    pdf = reportes.a_pdf(res.matriz, res.hallazgos, res.entidad)
    return "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)


def test_formato_unico():
    assert F.porcentaje(0.2229) == "22,3 %" and F.porcentaje_100(85.7) == "85,7 %"
    assert F.pesos(3884965483) == "$ 3.884.965.483" and F.numero(12000) == "12.000" and F.numero(41.667) == "41,667"
    assert F.numero(None) == F.pesos(float("nan")) == F.porcentaje(None) == "sin dato"


def test_pdf_trae_guia_de_lectura_y_marcas_por_meta(texto_pdf, res):
    assert "CÓMO LEER ESTE DOCUMENTO" in texto_pdf and "TODAS las actividades" in texto_pdf
    n = int((res.matriz["estado_reporte"] == "Reportada").sum())
    assert len(re.findall(r"INICIO DE LA META \d+ DE \d+", texto_pdf)) == n
    assert len(re.findall(r"FIN DE LA META MP", texto_pdf)) == n


def test_pdf_sin_fracciones_ni_formato_ingles_ni_simbolos_raros(texto_pdf):
    assert "■" not in texto_pdf and "ℹ" not in texto_pdf
    assert not re.search(r"\b0\.\d{3}\b", texto_pdf)                 # nada de 0.857 como avance
    assert not re.search(r"\d{1,3}(,\d{3}){2,}\.\d{2}", texto_pdf)   # nada de 20,901,759,592.00
    assert "Ejecución financiera (obligaciones / presupuesto definitivo): 60,0 %" in texto_pdf
    assert "$ 2.500.000.000" in texto_pdf


def test_excel_porcentajes_en_escala_0_100_y_hoja_leeme(res):
    wb = load_workbook(io.BytesIO(reportes.a_excel(res.matriz, res.hallazgos, res.calidad, res.aportes, res.entidad)))
    assert wb.sheetnames[0] == "Leeme"
    ws = wb["MP_PI_PA"]
    enc = [c.value for c in ws[1]]
    assert not any("(0-1)" in str(e) for e in enc)
    col = enc.index("Ejecución financiera (%)") + 1
    valores = [ws.cell(i, col).value for i in range(2, ws.max_row + 1) if ws.cell(i, col).value is not None]
    assert 60.0 in valores and all(0 <= v <= 100 for v in valores)


def test_prompt_explica_la_convencion_y_no_habla_de_decimales():
    p = generar_prompt_sistema("Revisión a Cierre de Vigencia", 2026, False)
    assert "CÓMO LEER LAS CIFRAS" in p and "0.2 =20%" not in p and "No los recalcules ni los multipliques por 100" in p
