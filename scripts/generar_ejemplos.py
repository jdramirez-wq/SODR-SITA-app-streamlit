"""Genera archivos de ejemplo FICTICIOS (entidad 9999) con la misma estructura y las mismas rarezas
de formato que los archivos reales de EVAPLAN y del libro de Drive.

Sirven para pruebas y documentación sin exponer datos reales (el repositorio es público).
Incluyen anomalías sembradas a propósito (ver docs/ANALISIS_FUENTES_EVAPLAN.md):
  - una meta existe en Drive pero falta en el export de EVAPLAN
  - una actividad con obligaciones y sin avance físico
  - una meta con PG incoherente en el PI original de Drive
  - Z023: una meta compartida (aportes de otra dependencia y de una entidad descentralizada sin código PS),
    una actividad de Centralizadas ausente del Z023, un BPIN no válido y una fila sin meta de producto

Uso:  python scripts/generar_ejemplos.py
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

DESTINO = Path(__file__).resolve().parent.parent / "data" / "ejemplos"
ENT, ENT_NOMBRE = "9999", "SECRETARÍA DE EJEMPLO"

# (código MP, MR, descripción, comportamiento, unidad, PG, original 2024-27, vigente 2024-27, resultado, proyectado)
METAS = [
    ("MP9900101019901001", "99001", "CAPACITAR 100 PERSONAS EN EJEMPLOS", "Incremento Acumulado", "Número",
     100, [20, 30, 30, 20], [20, 30, 30, 20], 12, 25),
    ("MP9900101029901002", "99001", "ALCANZAR 90 % DE DISPONIBILIDAD DEL SERVICIO", "Incremento Flujo", "Porcentaje",
     90, [70, 80, 85, 90], [70, 82.5, 85, 90], 41.667, None),
    ("MP9900202019902001", "99002", "MANTENER 100 % DE LA OPERACIÓN", "Mantenimiento Stock", "Número",
     100, [100, 100, 100, 100], [100, 100, 100, 100], 0, None),
    ("MP9900202029902002", "99002", "ENTREGAR 10 DOCUMENTOS TÉCNICOS", "Incremento Acumulado", "Número",
     10, [0, 4, 3, 3], ["NP", 4, 3, 3], 1, 3),
    # Solo en Drive (falta en el export de EVAPLAN) y con PG incoherente en el PI original:
    ("MP9900202039902003", "99002", "FORMULAR 1 PLAN DE EJEMPLO", "Incremento flujo", "Número",
     1, [0, 0, 1, 2], [0, 0, 1, 1], None, None),
]
# Otra entidad, solo para probar el filtro por entidad en Drive.
OTRA = ("MP9800101019801001", "98001", "META DE OTRA ENTIDAD", "Incremento Acumulado", "Número",
        50, [10, 10, 10, 20], [10, 10, 10, 20])

# (código proyecto, BPIN, nombre, MP, código actividad, nombre, fondo, inicial, definitivo, obligaciones, disponible,
#  estado, programada, ejecutada, % avance, observación)
ACTIVIDADES = [
    ("PI99-000001", 2024009990001, "Proyecto de ejemplo uno", "MP9900101019901001", "PI99-000001/1/1/01/01",
     "Realizar talleres", "2.000.000.000", "2.000.000.000", "1.500.000.000", "300.000.000", "CON EJECUCION", 40, 10, 25.0, "."),
    ("PI99-000001", 2024009990001, "Proyecto de ejemplo uno", "MP9900101019901001", "PI99-000001/1/1/01/02",
     "Elaborar material", "500.000.000", "500.000.000", 0, "500.000.000", "SIN EJECUCION", 5, None, 0.0, "."),
    ("PI99-000001", 2024009990001, "Proyecto de ejemplo uno", "MP9900101029901002", "PI99-000001/1/2/01/03",
     "Operar la plataforma", "1.200.000.000", "1.200.000.000", "800.000.000", "100.000.000", "CON EJECUCION", 12, 5, 41.67, "Avance según reporte"),
    ("PI99-000002", 2024009990002, "Proyecto de ejemplo dos", "MP9900202029902002", "PI99-000002/1/1/01/01",
     "Contratar consultoría", "900.000.000", "900.000.000", "450.000.000", "450.000.000", "CON EJECUCION", 4, None, 0.0, "."),
    ("PI99-000002", 2024009990002, "Proyecto de ejemplo dos", "MP9900202029902002", "PI99-000002/1/1/01/02",
     "Publicar documentos", "100.000.000", "100.000.000", 0, "100.000.000", "SIN EJECUCION", 6, None, 0.0, "."),
    # La MISMA actividad con un segundo registro presupuestal (ID distinto), como pasa en los datos reales.
    ("PI99-000002", 2024009990002, "Proyecto de ejemplo dos", "MP9900202029902002", "PI99-000002/1/1/01/01",
     "Contratar consultoría", "100.000.000", "100.000.000", 0, "100.000.000", "SIN EJECUCION", 4, None, 0.0, "."),
]


def _hoja(wb: Workbook, titulo: str | None = None):
    ws = wb.active
    if titulo:
        ws.title = titulo
    return ws


def _guardar(wb: Workbook, nombre: str) -> None:
    DESTINO.mkdir(parents=True, exist_ok=True)
    wb.save(DESTINO / nombre)


def _titulo_evaplan(ws, ncols: int) -> None:
    ws["A1"] = "EvaPlan"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)


def pi_mp_evaplan() -> None:
    enc = ["Codigo Entidad", "Nombre Entidad", "Código de Meta", "Descripción de Meta", "Codigo Programa",
           "Codigo Subprograma", "Comportamiento del Indicador", "Valor Linea Base", "Año Linea Base",
           "Periodicidad de Medición", "Unidad de Medida", "Variable", "Constante (K)", "Fórmula", "Resultado",
           "Valor Proyectado", "PG", 2024, 2025, 2026.0, 2027.0, "Principal Logro en Función del Cumplimiento",
           "Análisis del Logro", "Dificultades o Gestiones",
           "Negro, Mulato, Afrodescendiente, Raizal y Palenquero", "Indígena", "Otros", "¿Cuál Otro?", "Alcalá", "Cali"]
    wb = Workbook()
    ws = _hoja(wb, "Sheet1")
    _titulo_evaplan(ws, len(enc))
    ws.append(enc)
    for mp, mr, desc, comp, unidad, pg, _orig, vig, res, proy in METAS[:4]:
        sub = mp[7:9]
        ws.append([int(ENT), ENT_NOMBRE, mp, desc, "99 - Programa de Ejemplo", f"{sub} - Subprograma Uno",
                   comp.replace("flujo", "Flujo"), 0, 2023, "Anual", unidad, "Nombre variable: V de ejemplo",
                   0, "V1", res, proy, pg, *vig, "Se está cumpliendo" if res else None, "Análisis de ejemplo",
                   "Dificultad de ejemplo", None, None, 0, "Servidores públicos", 0, 0])
    _guardar(wb, "ejemplo_PI_MP_evaplan.xlsx")


def pi_mr_evaplan() -> None:
    enc = ["Codigo Entidad", "Nombre Entidad", "Código de Meta", "Descripción de Meta", "Codigo Programa",
           "Comportamiento del Indicador", "Valor Linea Base", "Año Linea Base", "Periodicidad de Medición",
           "Unidad de Medida", "Variable", "Constante (K)", "Fórmula", "Resultado", "Valor Proyectado", "PG",
           2024, 2025, 2026, 2027, "Principal Logro en Función del Cumplimiento", "Análisis del Logro",
           "Dificultades o Gestiones"]
    wb = Workbook()
    ws = _hoja(wb, "Sheet1")
    _titulo_evaplan(ws, len(enc))
    ws.append(enc)
    for cod, desc, res in (("99001", "ALCANZAR 40 PUNTOS EN EL ÍNDICE DE EJEMPLO", 0), ("99002", "ALCANZAR 22,7 % EN EL INDICADOR B", 21.7)):
        ws.append([int(ENT), ENT_NOMBRE, int(cod), desc, "99 - Programa de Ejemplo", "Incremento Flujo", 25, 2022,
                   "Anual", "Puntos", "Nombre variable: puntaje", 0, "V1", res, None, 40, 50, 30, 35, 40, ".", ".",
                   "Se mide a final de año"])
    _guardar(wb, "ejemplo_PI_MR_evaplan.xlsx")


def centralizadas() -> None:
    enc = ["ID", "Cod.Entidad", "Nombre.Entidad", "Cód. Proyecto", "Nombre Proyecto", "Cód. BPIN", "Cód. MP",
           "Descripción MP", "Cód.Producto MGA", "Producto MGA", "Cod.Indicador Producto MGA",
           "Indicador Producto MGA", "Cód. Actividad", "Nombre Actividad", "Cód. Fondo", "Nombre Fondo",
           "Ppto. Inicial", "Ppto. Definitivo", "Ppto. Total Obligaciones", "Ppto. Disponible", "Ppto. Gestión",
           "Cód. Fondo Gestión", "Fondo Gestión", "Estado Actividad", "Unid. Medida", "Complemento",
           "Unid. + Compl.", "Cant. Prog. Vig.", "Cant. Ejec. Vig.", "% Avance x Actividad", "Observación", "Estado"]
    wb = Workbook()
    ws = _hoja(wb, "Sheet1")
    _titulo_evaplan(ws, len(enc))
    ws.append(enc)
    desc = {m[0]: m[2] for m in METAS}
    for i, (proy, bpin, nproy, mp, act, nact, ini, defi, obl, disp, est, prog, ejec, pct, obs) in enumerate(ACTIVIDADES):
        prod = mp[-7:]
        ws.append([12000 + i, int(ENT), "SRIA DE EJEMPLO  DE LA INFORM", proy, nproy, bpin, mp, desc[mp], int(prod),
                   f"Servicio de ejemplo {prod}", int(prod + "00"), f"Personas atendidas {prod}", act, nact,
                   121000, "Ingresos corrientes de Libre Destinación", ini, defi, obl, disp, None, None, None, est,
                   "Número", "de personas", "Número de personas", prog, ejec, pct, obs, "ACTIVO"])
    _guardar(wb, "ejemplo_Centralizadas.xlsx")


def _bandas(ws, bandas: list[tuple[int, int, str]]) -> None:
    for ini, fin, texto in bandas:
        ws.cell(1, ini + 1, texto)
        if fin > ini:
            ws.merge_cells(start_row=1, start_column=ini + 1, end_row=1, end_column=fin + 1)


def drive() -> None:
    wb = Workbook()
    # ---------------- hoja MP (103 columnas leídas por el lector)
    ws = wb.active
    ws.title = "MP"
    enc = [None] * 103
    base = ["Código MP", "Meta Producto", "Entidad", "Línea Estrategica", "Línea Programa", "Meta Resultado Asociada",
            "Subprograma", "Indicador Principal", "Cód - Objetivos de Desarrollo Sostenible - ODS",
            "Descrip - Objetivos de Desarrollo Sostenible - ODS", "Cód Meta ODS", "Descripción Meta ODS", "Sector MGA",
            "Sector SAP", "Sector Nombre", "Programa Catálogo", "Producto Catálogo", "Descripcion del producto",
            "Indicador de Producto", "Indicador de Resultado", "Area Funcional", "Unidad de medida", "Comportamiento",
            "Valor Linea Base", "Año Linea Base", "Periodicidad de medición", "Fecha Actualizacion Indicador",
            "Meta PG", "VARIABLES", "FÓRMULA", "PG 2024-2027.", "2024.", "2025.", "2026.", "2027.", "PG 2024-2027",
            "VAL ALC 2024\n", "VAL ALC 2025", 2026, 2027,
            "VERIFICACIÓN DEL COMPORTAMIENTO DE META EN REPROGRAMACIÓN", "EDT\n(corte 02/03/2026)",
            "ESTADO DE LA META DE PRODUCTO"]
    # el real tiene 'VAL ALC 2024', 'VAL ALC 2025', '2026', '2027' en 36-39; 'VERIF…' en 40, 'EDT' en 41, 'ESTADO' en 42
    base = base[:36] + ["VAL ALC 2024\n", "VAL ALC 2025", 2026, 2027] + base[40:]
    enc[:len(base)] = base
    fuentes = ["ICLD", "DEST. ESP", "SGP EDUCACIÓN", "SGP SALUD", "SGP APSB", "SGR ASIG. REGIONAL", "SGR CT&I",
               "CONFINANCIACIÓN NACIÓN", "CRÉDITO", "GESTIÓN", "PROPIOS DESCENTRALIZADOS"]
    for bloque, anio in enumerate(("2024", "2025", "2026", "2027", "2024-2027")):
        ini = 43 + bloque * 12
        for j, f in enumerate(fuentes):
            enc[ini + j] = f"{f} {anio}"
        enc[ini + 11] = f"TOTAL RECURSO {anio}"
    for i, h in enumerate(enc):
        if h is None:
            enc[i] = f"col_{i}"
    _bandas(ws, [(0, 2, "INFORMACIÓN PDD 2024-2027"), (30, 34, "PLAN INDICATIVO - PI"),
                 (37, 39, "REPROGRAMACIÓN PLAN DE ACCIÓN - PA"), (43, 54, "RECURSOS 2024"), (55, 66, "RECURSOS 2025"),
                 (67, 78, "RECURSOS 2026"), (79, 90, "RECURSOS 2027"), (91, 102, "RECURSOS PG")])
    ws.append([])  # fila 1 ya tiene las bandas; reescribimos encabezados en la fila 2
    for c, h in enumerate(enc, start=1):
        ws.cell(2, c, h)

    def fila(mp, mr, desc, comp, unidad, pg, orig, vig, ent_cod, ent_nom, recurso):
        f = [None] * 103
        sub = mp[7:9]
        f[0], f[1] = mp, f"{mp} - {desc}"
        f[2] = f"{ent_cod} - {ent_nom}"
        f[3], f[4] = "1 - Línea de Ejemplo", f"{mp[2:4]} - Programa de Ejemplo"
        f[5] = f"{mr} - META DE RESULTADO {mr}"
        f[6] = f"{sub} - Subprograma Uno"
        f[21], f[22] = unidad, comp
        f[23], f[24] = ("NO DISPONIBLE", "2020-2023") if mp.endswith("2003") else (0, 2023)
        f[25], f[26] = "ANUAL", "2024-12-12"
        f[27], f[28], f[29] = pg, "V1 = VARIABLE DE EJEMPLO", "V1"
        f[30], f[31:35] = pg, orig
        f[35], f[36:40] = pg, vig
        f[40] = {"Incremento Acumulado": "ACUMULADO", "Incremento Flujo": "FLUJO",
                 "Incremento flujo": "FLUJO", "Mantenimiento Stock": "MANTENIMIENTO"}[comp]
        f[41], f[42] = ("Sí" if mp.endswith("2001") else "No"), "EN EJECUCIÓN"
        f[54], f[66], f[78], f[90], f[102] = recurso / 4, recurso / 4, recurso / 4, recurso / 4, recurso
        return f

    for mp, mr, desc, comp, unidad, pg, orig, vig, *_ in METAS:
        ws.append(fila(mp, mr, desc, comp, unidad, pg, orig, vig, ENT, ENT_NOMBRE, 12_000_000_000.0))
    ws.append(fila(*OTRA[:6], OTRA[6], OTRA[7], "9998", "OTRA SECRETARÍA DE EJEMPLO", 1_000_000_000.0))
    ws.append([])  # filas vacías al final, como en el libro real
    ws.append([])

    # ---------------- hoja MR
    mr = wb.create_sheet("MR")
    enc_mr = ["COD MR", "Meta resultado", "Entidad", "Nombre del Indicador", "Linea Estratégica", "Programa Plan",
              "Fuente Indicador", "Categoria Terridata", "Tranformación Pnd", "Trazador Asociado", "Sigla Trazador",
              "Descripcion del Indicador", "Unidad de medida", "Valor linea Base", "Año linea Base",
              "Comportamiento del indicador", "Meta PG", "Periodicidad de medición", "Fecha Actualización Indicador",
              "VARIABLES", "FORMULA", "PG", 2024, 2025, 2026, 2027, "PG", "VAL ALC 2024\n", 2025, 2026, 2027,
              "VERIFICACIÓN DEL COMPORTAMIENTO DE META EN REPROGRAMACIÓN"]
    _bandas(mr, [(21, 25, "PROGRAMACIÓN ORIGINAL"), (26, 30, "LOGRO REPROGRAMACIÓN")])
    for c, h in enumerate(enc_mr, start=1):
        mr.cell(2, c, h)
    filas_mr = [("MR99001", "99001", "ALCANZAR 40 PUNTOS EN EL ÍNDICE DE EJEMPLO", ENT, ENT_NOMBRE, "Semetral", "FLUJO"),
                ("MR99002", "99002", "ALCANZAR 22,7 % EN EL INDICADOR B", ENT, ENT_NOMBRE, "Anual", "ERROR"),
                ("MR98001", "98001", "META DE RESULTADO DE OTRA ENTIDAD", "9998", "OTRA SECRETARÍA DE EJEMPLO", "Anual", "FLUJO")]
    for cod, num, desc, ec, en, perio, verif in filas_mr:
        f = [None] * 32
        f[0], f[1] = cod, f"{num} - {num}-{desc}"
        f[2], f[3], f[4], f[5] = f"{ec} - {en}", "Indicador de ejemplo", "1 - Línea de Ejemplo", "99 - Programa de Ejemplo"
        f[12], f[13], f[14], f[15], f[16] = "Puntos", 25, 2022, "Reducción Anual" if cod == "MR98001" else "Incremento Flujo", 40
        f[17], f[18] = perio, "2024-08-01"
        f[19], f[20] = "V1 = puntaje", "V1"
        f[21:26] = [40, 50, 30, 35, 40]
        f[26:31] = [40, 50, 30, 35, 40]
        f[31] = verif
        mr.append(f)
    # la hoja MR real tiene las filas de datos desde la fila 3: ya quedaron en 3.. (append tras fila 2)
    _guardar(wb, "ejemplo_PI_Drive.xlsx")


Z023_ENCABEZADOS = [
    "Dependencia", "Descripción Dependencia", "PPM: Proyecto", "Código Programa", "Descripción Programa",
    "Descripción PROYECTO", "Sector", "Descripción Sector", "Area Funcional", "Cod.BPIN DNP", "Subprograma",
    "Descripción de Subprograma", "Programa", "Descripción de Programa", "Linea de Accion",
    "Descripción de Linea de Accion", "Pilar / LT", "Descripción de Pilar / LT", "Plan de Desarrollo",
    "Descripción de Plan de Desarrollo", "Codigo Meta Resultado", "Descripción de la Meta Resultado",
    "Codigo Meta Producto", "Descripción de la Meta Producto", "PPM: Producto", "PS: Producto",
    "Código Producto DNP", "Descripción Producto DNP", "PPM: Actividad", "PS: Actividad", "Descripción Actividad",
    "Pos.Pre.", "Vigencia", "Fondo", "Descripción Fondo", "Recurso", "Tipo", "Tipo Descripcion", "Tipo Proyecto",
    "Tipo Proyecto Descripcion", "Entidad", "Entidad Descripcion", "Participacion Ciudadana", "Meta Actividad",
    "Gasto Social", "Gasto FBK/GO", "Tipo Actividad", "Pasivo exigible vigencia expirada", " Valor de Actividad",
    "Politica Publica Linea Estrategica", "Trazador 01", "Trazador 02", "Trazador 03", "Trazador 04", "Trazador 05",
    "Pres. Part.", "Sector DNP", "SECTOR MGA", "PROGRAMA MGA", "PRODUCTO ASOCIADO A LA MP",
    "Fecha de importación de los datos", "Centro Gestor de la MP", "PRODUCTO DIRECTO AL PDD?", "ACUMULA",
    "INDICADOR", "LLAVE DNP - PIIP", "LLAVE DNP - CUIPO", "¿EDT?", "ID MGA", "Objetivo General Proyecto",
    "Objetivo Específico", "Columna1", "Columna2",
]

# (dependencia, nombre, PPM proyecto, nombre proyecto, BPIN, MP, PPM actividad, PS actividad, nombre actividad,
#  vigencia, valor)
Z023_FILAS = [
    ("9999", "SRIA DE EJEMPLO", "PI-900001", "Proyecto de ejemplo uno", "2024009990001", "MP9900101019901001",
     "0000000001", "PI99-000001/1/1/01/01", "Realizar talleres", 2026, 2000000000),
    ("9999", "SRIA DE EJEMPLO", "PI-900001", "Proyecto de ejemplo uno", "2024009990001", "MP9900101019901001",
     "0000000002", "PI99-000001/1/1/01/02", "Elaborar material", 2026, 500000000),
    ("9999", "SRIA DE EJEMPLO", "PI-900001", "Proyecto de ejemplo uno", "2024009990001", "MP9900101029901002",
     "0000000003", "PI99-000001/1/2/01/03", "Operar la plataforma", 2026, 1200000000),
    ("9999", "SRIA DE EJEMPLO", "PI-900002", "Proyecto de ejemplo dos", "2024009990002", "MP9900202029902002",
     "0000000004", "PI99-000002/1/1/01/01", "Contratar consultoría", 2026, 1000000000),
    # (falta a propósito PI99-000002/1/1/01/02: está en Centralizadas pero no en el Z023)
    # Misma actividad en una vigencia anterior: debe ignorarse al filtrar por vigencia.
    ("9999", "SRIA DE EJEMPLO", "PI-900001", "Proyecto de ejemplo uno", "2024009990001", "MP9900101019901001",
     "0000000005", "PI99-000001/1/1/01/01", "Realizar talleres", 2025, 900000000),
    # Meta compartida: aporta otra dependencia (con código PS) y una entidad descentralizada (sin código PS).
    ("9998", "OTRA SRIA DE EJEMPLO", "PI-900003", "Proyecto de otra dependencia", "2025009980003",
     "MP9900101019901001", "0000000006", "PI98-000003/1/1/01/01", "Apoyar los talleres", 2026, 300000000),
    ("0099", "ENTIDAD DESCENTRALIZADA DE EJEMPLO", "PI-900004", "Proyecto de la entidad descentralizada", "202400099004",
     "MP9900101019901001", "0000000007", None, "Capacitar a su personal", 2026, 150000000),
    # Anomalías: BPIN no válido (7 dígitos) y actividad sin meta de producto.
    ("9998", "OTRA SRIA DE EJEMPLO", "PI-900005", "Proyecto con BPIN malo", "1234567", "MP9800101019801001",
     "0000000008", "PI98-000005/1/1/01/01", "Actividad de otra entidad", 2026, 100000000),
    ("9998", "OTRA SRIA DE EJEMPLO", "PI-900005", "Proyecto con BPIN malo", "2025009980005", None,
     "0000000009", "PI98-000005/1/1/01/02", "Actividad sin meta", 2026, 50000000),
]


def z023() -> None:
    """Z023 consolidado ficticio: hoja 'Hoja1' con las 73 columnas reales; solo se llenan las que lee el esquema."""
    wb = Workbook()
    ws = _hoja(wb, "Menú")
    ws["A1"] = "CONSOLIDAR Z023"
    hoja = wb.create_sheet("Hoja1")
    hoja.append(Z023_ENCABEZADOS)
    idx = {h.strip(): i for i, h in enumerate(Z023_ENCABEZADOS)}
    for dep, nom, ppm, nproy, bpin, mp, ppm_act, ps_act, nact, vig, valor in Z023_FILAS:
        f = [None] * len(Z023_ENCABEZADOS)
        f[idx["Dependencia"]], f[idx["Descripción Dependencia"]] = dep, nom
        f[idx["PPM: Proyecto"]], f[idx["Descripción PROYECTO"]], f[idx["Cod.BPIN DNP"]] = ppm, nproy, bpin
        f[idx["Codigo Meta Resultado"]] = ("MR" + mp[2:7]) if mp else None
        f[idx["Codigo Meta Producto"]] = mp
        f[idx["PS: Producto"]] = ps_act.rsplit("/", 2)[0].replace("/", "") if ps_act else None
        f[idx["PPM: Actividad"]], f[idx["PS: Actividad"]] = ppm_act, ps_act
        f[idx["Descripción Actividad"]], f[idx["Vigencia"]] = nact, str(vig)
        f[idx["Fondo"]], f[idx["Descripción Fondo"]] = "121000", "Ingresos corrientes de Libre Destinación"
        f[idx["Tipo Actividad"]], f[idx["Valor de Actividad"]] = "Inversión", valor
        f[idx["Centro Gestor de la MP"]] = "#ERROR!"        # columna de fórmula que a veces llega con error
        hoja.append(f)
    _guardar(wb, "ejemplo_Z023.xlsx")


if __name__ == "__main__":
    z023()
    pi_mp_evaplan()
    pi_mr_evaplan()
    centralizadas()
    drive()
    print(f"Ejemplos generados en {DESTINO}")
